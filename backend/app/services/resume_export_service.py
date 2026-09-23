import io
from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.models.profile import CandidateProfile, CuratedResume

class ResumeExportService:
    @staticmethod
    def generate_pdf(profile: CandidateProfile, curated: CuratedResume) -> bytes:
        """
        Generates a clean, ATS-compliant, single-column PDF resume.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=36,
            rightMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'DocTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=18,
            leading=22,
            alignment=TA_CENTER,
            textColor=colors.HexColor('#1E293B')
        )

        subtitle_style = ParagraphStyle(
            'DocSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=11,
            leading=15,
            alignment=TA_CENTER,
            textColor=colors.HexColor('#475569')
        )

        contact_style = ParagraphStyle(
            'Contact',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            alignment=TA_CENTER,
            textColor=colors.HexColor('#64748B')
        )

        section_heading = ParagraphStyle(
            'SectionHeading',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=12,
            leading=16,
            alignment=TA_LEFT,
            textColor=colors.HexColor('#0F172A'),
            spaceBefore=8,
            spaceAfter=3
        )

        role_heading = ParagraphStyle(
            'RoleHeading',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#1E293B')
        )

        company_meta = ParagraphStyle(
            'CompanyMeta',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor('#475569')
        )

        body_style = ParagraphStyle(
            'Body',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor('#334155'),
            spaceAfter=2
        )

        bullet_style = ParagraphStyle(
            'Bullet',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            leftIndent=14,
            textColor=colors.HexColor('#334155'),
            spaceAfter=2
        )

        story = []

        # Name and Headline
        story.append(Paragraph(f"{profile.first_name} {profile.last_name}", title_style))
        story.append(Spacer(1, 2))
        story.append(Paragraph(curated.tailored_headline, subtitle_style))
        story.append(Spacer(1, 4))

        # Contact Info
        contact_line = f"{profile.email} | {profile.phone} | {profile.location}"
        if profile.linkedin_url:
            contact_line += f" | {profile.linkedin_url.replace('https://', '')}"
        if profile.github_url:
            contact_line += f" | {profile.github_url.replace('https://', '')}"
        story.append(Paragraph(contact_line, contact_style))
        story.append(Spacer(1, 6))
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=6))

        # Executive Summary
        if curated.tailored_summary:
            story.append(Paragraph("PROFESSIONAL SUMMARY", section_heading))
            story.append(Paragraph(curated.tailored_summary, body_style))
            story.append(Spacer(1, 6))

        # Highlighted Skills
        if curated.highlighted_skills:
            story.append(Paragraph("TECHNICAL & CORE SKILLS", section_heading))
            skills_text = " • ".join(curated.highlighted_skills)
            story.append(Paragraph(skills_text, body_style))
            story.append(Spacer(1, 6))

        # Work Experience
        if curated.tailored_experiences:
            story.append(Paragraph("PROFESSIONAL EXPERIENCE", section_heading))
            for exp in curated.tailored_experiences:
                exp_title = f"{exp.role} — <b>{exp.company}</b>"
                date_str = f"{exp.start_date} – {exp.end_date}"
                if exp.location:
                    date_str = f"{exp.location} | {date_str}"
                story.append(Paragraph(exp_title, role_heading))
                story.append(Paragraph(date_str, company_meta))
                for bp in exp.bullet_points:
                    story.append(Paragraph(f"• {bp}", bullet_style))
                story.append(Spacer(1, 4))

        # Education
        if profile.education:
            story.append(Paragraph("EDUCATION", section_heading))
            for edu in profile.education:
                edu_line = f"<b>{edu.institution}</b> — {edu.degree}"
                if edu.field_of_study:
                    edu_line += f" ({edu.field_of_study})"
                if edu.start_year or edu.end_year:
                    edu_line += f" | {edu.start_year} - {edu.end_year}"
                story.append(Paragraph(edu_line, body_style))
            story.append(Spacer(1, 4))

        # Projects
        if curated.tailored_projects:
            story.append(Paragraph("PROJECTS", section_heading))
            for proj in curated.tailored_projects:
                proj_title = f"<b>{proj.title}</b>"
                if proj.technologies:
                    proj_title += f" [<i>{', '.join(proj.technologies)}</i>]"
                story.append(Paragraph(proj_title, role_heading))
                story.append(Paragraph(proj.description, body_style))
                if proj.link:
                    story.append(Paragraph(f"Link: {proj.link}", company_meta))
                story.append(Spacer(1, 3))

        doc.build(story)
        return buffer.getvalue()

    @staticmethod
    def generate_docx(profile: CandidateProfile, curated: CuratedResume) -> bytes:
        """
        Generates a clean ATS-friendly Microsoft Word DOCX resume.
        """
        doc = docx.Document()

        # Adjust margins to 0.5 in
        for section in doc.sections:
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

        # Name
        p_name = doc.add_paragraph()
        p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_name = p_name.add_run(f"{profile.first_name} {profile.last_name}")
        run_name.font.size = Pt(18)
        run_name.font.bold = True
        run_name.font.name = 'Calibri'

        # Headline
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_sub = p_sub.add_run(curated.tailored_headline)
        run_sub.font.size = Pt(11)
        run_sub.font.color.rgb = RGBColor(71, 85, 105)

        # Contact
        p_contact = doc.add_paragraph()
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_line = f"{profile.email} | {profile.phone} | {profile.location}"
        if profile.linkedin_url:
            contact_line += f" | {profile.linkedin_url}"
        if profile.github_url:
            contact_line += f" | {profile.github_url}"
        run_contact = p_contact.add_run(contact_line)
        run_contact.font.size = Pt(9)
        run_contact.font.color.rgb = RGBColor(100, 116, 139)

        def add_section_header(title: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(title)
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(15, 23, 42)

        # Summary
        if curated.tailored_summary:
            add_section_header("PROFESSIONAL SUMMARY")
            p = doc.add_paragraph(curated.tailored_summary)
            p.paragraph_format.space_after = Pt(4)

        # Skills
        if curated.highlighted_skills:
            add_section_header("TECHNICAL SKILLS")
            p = doc.add_paragraph(" • ".join(curated.highlighted_skills))
            p.paragraph_format.space_after = Pt(4)

        # Experience
        if curated.tailored_experiences:
            add_section_header("PROFESSIONAL EXPERIENCE")
            for exp in curated.tailored_experiences:
                p_exp = doc.add_paragraph()
                r_role = p_exp.add_run(f"{exp.role} — {exp.company}")
                r_role.font.bold = True
                date_str = f" ({exp.start_date} – {exp.end_date})"
                r_dates = p_exp.add_run(date_str)
                r_dates.font.italic = True
                p_exp.paragraph_format.space_after = Pt(1)

                for bp in exp.bullet_points:
                    bp_p = doc.add_paragraph(f"{bp}", style='List Bullet')
                    bp_p.paragraph_format.space_after = Pt(1)

        # Education
        if profile.education:
            add_section_header("EDUCATION")
            for edu in profile.education:
                p = doc.add_paragraph()
                r = p.add_run(f"{edu.institution} — {edu.degree}")
                r.font.bold = True
                if edu.field_of_study:
                    p.add_run(f", {edu.field_of_study}")
                if edu.end_year:
                    p.add_run(f" ({edu.end_year})")

        # Projects
        if curated.tailored_projects:
            add_section_header("KEY PROJECTS")
            for proj in curated.tailored_projects:
                p = doc.add_paragraph()
                r = p.add_run(proj.title)
                r.font.bold = True
                if proj.technologies:
                    p.add_run(f" [{', '.join(proj.technologies)}]")
                doc.add_paragraph(proj.description)

        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()
