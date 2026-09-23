import pytest
import io
from fastapi.testclient import TestClient
from docx import Document
from PIL import Image

from app.main import app
from app.services.parser_service import ResumeParserService
from app.services.profile_store import ProfileStore
from app.models.profile import CandidateProfile, JobAnalysisResult, CuratedResume, WorkExperience

client = TestClient(app)

def test_health_check():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "online"

def test_dashboard_and_demo_endpoints():
    res_dash = client.get("/dashboard")
    assert res_dash.status_code == 200
    assert "ADONIS" in res_dash.text

    res_demo = client.get("/demo")
    assert res_demo.status_code == 200
    assert "application_form" in res_demo.text

def test_get_and_update_profile():
    # Get initial profile
    res = client.get("/api/profile")
    assert res.status_code == 200
    profile = res.json()
    assert "email" in profile
    assert profile["first_name"] == "Alex"

    # Update profile
    profile["headline"] = "Principal AI & Systems Engineer"
    update_res = client.put("/api/profile", json=profile)
    assert update_res.status_code == 200
    assert update_res.json()["headline"] == "Principal AI & Systems Engineer"

def test_docx_parsing():
    doc = Document()
    doc.add_heading("John Doe", level=0)
    doc.add_paragraph("john.doe@example.com | 123-456-7890")
    doc.add_paragraph("Full stack engineer with Python and React skills.")
    buffer = io.BytesIO()
    doc.save(buffer)
    docx_bytes = buffer.getvalue()

    text = ResumeParserService.extract_text_from_docx(docx_bytes)
    assert "John Doe" in text
    assert "john.doe@example.com" in text
    assert "Python and React" in text

def test_pdf_parsing():
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(100, 750, "Jane Smith")
    c.drawString(100, 730, "jane.smith@example.com | 555-987-6543")
    c.drawString(100, 710, "Senior Cloud Engineer with AWS, Kubernetes, and Go experience.")
    c.save()
    pdf_bytes = buf.getvalue()

    text = ResumeParserService.extract_text_from_pdf(pdf_bytes)
    assert "Jane Smith" in text
    assert "jane.smith@example.com" in text
    assert "Cloud Engineer" in text

def test_image_parsing():
    img = Image.new("RGB", (200, 100), color=(255, 255, 255))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    img_bytes = buf.getvalue()

    mime, b64 = ResumeParserService.process_image(img_bytes)
    assert mime == "image/png"
    assert len(b64) > 0

def test_job_analysis_and_curation():
    sample_job_text = """
    Senior Full Stack Engineer at Stripe.
    We are looking for an experienced developer with deep knowledge of React, TypeScript, Python, and REST APIs.
    You will lead architectural design, optimize database queries in PostgreSQL, and build delightful user experiences.
    Requirements:
    - 4+ years of React and TypeScript experience
    - Proficiency in Python, FastAPI, or Node.js
    - Experience with Docker and cloud deployments
    """
    
    # Analyze job
    analysis_res = client.post("/api/job/analyze", json={
        "job_title": "Senior Full Stack Engineer",
        "company_name": "Stripe",
        "raw_text": sample_job_text
    })
    assert analysis_res.status_code == 200
    job_analysis = analysis_res.json()
    assert job_analysis["company_name"] == "Stripe"
    assert "React" in job_analysis["required_skills"] or "TypeScript" in job_analysis["required_skills"]

    # Curate resume
    curate_res = client.post("/api/resume/curate", json={
        "job_analysis": job_analysis
    })
    assert curate_res.status_code == 200
    curated = curate_res.json()
    assert curated["ats_match_score"] >= 50
    assert len(curated["highlighted_skills"]) > 0
    assert "Stripe" in curated["suggested_cover_letter"] or "Senior Full Stack Engineer" in curated["tailored_headline"]

def test_screening_qa_generation():
    res = client.post("/api/screening/answers", json={
        "job_title": "Staff Engineer",
        "company_name": "OpenAI",
        "questions": [
            "Are you legally authorized to work in the country?",
            "What are your salary expectations?",
            "Why do you want to join OpenAI?"
        ]
    })
    assert res.status_code == 200
    answers = res.json()["answers"]
    assert len(answers) == 3

def test_pdf_and_docx_export():
    profile = ProfileStore.get_profile()
    curated = CuratedResume(
        tailored_headline="Senior Full Stack Engineer",
        tailored_summary="Expert engineer in React and Python.",
        highlighted_skills=["React", "TypeScript", "Python", "FastAPI"],
        tailored_experiences=profile.experiences,
        tailored_projects=profile.projects,
        ats_match_score=88,
        matching_keywords=["React", "Python"],
        missing_keywords=[],
        suggested_cover_letter="Sample cover letter"
    )

    # Test PDF export
    pdf_res = client.post("/api/resume/export/pdf", json=curated.model_dump())
    assert pdf_res.status_code == 200
    assert pdf_res.headers["content-type"] == "application/pdf"
    assert len(pdf_res.content) > 1000

    # Test DOCX export
    docx_res = client.post("/api/resume/export/docx", json=curated.model_dump())
    assert docx_res.status_code == 200
    assert "officedocument" in docx_res.headers["content-type"]
    assert len(docx_res.content) > 1000
